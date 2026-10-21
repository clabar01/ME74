"""Consolidate the field-test logs into one workbook (one tab per trial + all trials together).

Day 1 (2026-10-15): ~/Downloads/Data.zip  -> ESP-NOW 0-80 m walk + one IMU walk.
Day 2 (2026-10-20): ~/Downloads/Data 2/   -> dashboard exports (zips + loose overview CSVs), GPS/LoRa/ESP-NOW/IMU.

Run:  python3 analysis/build_workbook.py [downloads_dir]
Writes ME74_field_tests_consolidated.xlsx (+ _all_trials.csv) to the downloads folder. Source files are only read.
"""
import glob, io, os, re, sys, zipfile
import pandas as pd

DL = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Downloads")
OUT = os.path.join(DL, "ME74_field_tests_consolidated.xlsx")
EDT = pd.Timedelta(hours=4)

# ---------------------------------------------------------------- day 2
def read_day2_overviews():
    frames = []
    for z in glob.glob(os.path.join(DL, "Data 2", "*.zip")):
        with zipfile.ZipFile(z) as zf:
            for n in zf.namelist():
                if "overview" in n and n.endswith(".csv"):
                    frames.append(pd.read_csv(io.BytesIO(zf.read(n))))
    for f in glob.glob(os.path.join(DL, "Data 2", "ME74_overview_*.csv")):
        frames.append(pd.read_csv(f))
    # every export repeats the earlier rows of its session
    return pd.concat(frames).drop_duplicates("pc_time").sort_values("pc_time").reset_index(drop=True)

T = pd.Timestamp
DIR = {"Trial 1": "0->90", "Return run": "90->0", "Trial 2": "0->90", "Trial 3": "0->90", "Trial 4": "90->0"}

def day2_trial(t, lab):
    lab = str(lab)
    if t < T("2026-10-20T17:52:00Z"): return "Trial 1"
    if t < T("2026-10-20T18:00:00Z"): return "Return run"
    if t < T("2026-10-20T18:33:50Z"): return "Trial 2"            # aborted run
    if "trial 4" in lab.lower() or (t > T("2026-10-20T18:38:52Z") and "Trial 3" not in lab): return "Trial 4"
    return "Trial 3"

def clean_label(lab, trial):
    """Full segment label typed for this trial, or None while it is being typed / left over / a typo."""
    if not isinstance(lab, str): return None
    m = re.fullmatch(r"\s*(\d+)-(\d+)\s*m?\s*(trial\s*(\d))?\s*", lab, re.I)
    if not m: return None
    a, b = int(m.group(1)), int(m.group(2))
    if a % 10 or b % 10 or abs(a - b) != 10 or max(a, b) > 90: return None
    if m.group(4) and f"Trial {m.group(4)}" != trial: return None
    if not m.group(4) and trial in ("Trial 2", "Trial 3", "Trial 4"): return None
    if (a < b) != (DIR[trial] == "0->90"): return None
    return f"{a}-{b}"

d2 = read_day2_overviews()
d2["t"] = pd.to_datetime(d2.pc_time)
d2["Trial"] = [day2_trial(t, l) for t, l in zip(d2.t, d2.test)]
d2 = d2[d2.Trial != "Trial 2"].copy()                          # aborted run, left out
d2["Segment (m)"] = [clean_label(l, t) for l, t in zip(d2.test, d2.Trial)]
# rows logged while a label was typed (or before the first one) belong to the next segment
d2["Segment (m)"] = d2.groupby("Trial")["Segment (m)"].transform(lambda s: s.bfill().ffill())
# drop rows logged before the sensors were zeroed at the start of a run (stale from-zero values)
first_zero = d2[d2.event.notna()].groupby("Trial").t.min()
d2 = d2[~d2.apply(lambda r: r.Trial in first_zero and r.t < first_zero[r.Trial], axis=1)]

COLS = {"event": "Event", "gps_from_zero_m": "GPS dist from zero (m)", "gps_sats": "GPS sats", "gps_hdop": "GPS HDOP",
        "lora_dist_m": "LoRa est dist (m)", "lora_since_zero_m": "LoRa since zero (m)", "lora_rssi_dbm": "LoRa RSSI (dBm)",
        "lora_snr_db": "LoRa SNR (dB)", "espnow_dist_m": "ESP-NOW est dist (m)", "espnow_since_zero_m": "ESP-NOW since zero (m)",
        "espnow_rssi_dbm": "ESP-NOW RSSI (dBm)", "espnow_loss_pct": "ESP-NOW loss (%)",
        "imu_from_zero_m": "IMU1 dist from zero (m)", "imu2_from_zero_m": "IMU2 dist from zero (m)"}
d2 = d2.rename(columns=COLS)
d2["Day"] = "Day 2 (10/20)"
d2["Event"] = d2["Event"].str.replace("zero_", "zero ").str.replace(";", ", ")

# ---------------------------------------------------------------- day 1
with zipfile.ZipFile(os.path.join(DL, "Data.zip")) as zf:
    esp = []
    for n in zf.namelist():
        m = re.search(r"espnow_(\d+)-(\d+)m(\(\d\))?\.csv$", n)
        if m:
            f = pd.read_csv(io.BytesIO(zf.read(n)))
            f["Segment (m)"] = f"{m.group(1)}-{m.group(2)}"
            esp.append(f)
    imu = pd.read_csv(io.BytesIO(zf.read("Data/IMU Data.csv")))
d1 = pd.concat(esp).drop_duplicates("time_iso").sort_values("time_iso")
d1["t"] = pd.to_datetime(d1.time_iso)
d1 = d1.rename(columns={"est_distance_m": "ESP-NOW est dist (m)", "rssi_smoothed_dbm": "ESP-NOW RSSI (dBm)", "packet_loss_pct": "ESP-NOW loss (%)"})
d1["Event"] = d1.status.map(lambda s: "ESP-NOW link lost" if s == "LOST" else None)
d1["Trial"], d1["Day"] = "ESP-NOW walk", "Day 1 (10/15)"
DIR["ESP-NOW walk"] = "0->90"

# ---------------------------------------------------------------- combine
allr = pd.concat([d1, d2], ignore_index=True)
allr["Time (s)"] = (allr.t - allr.groupby("Trial").t.transform("min")).dt.total_seconds().round(1)
allr["Clock (EDT)"] = (allr.t - EDT).dt.strftime("%H:%M:%S")
allr["Direction"] = allr.Trial.map(DIR)
ordered = lambda s: "-".join(map(str, sorted(map(int, s.split("-")))))
allr["Segment (0-90 order)"] = allr["Segment (m)"].map(ordered)
KEEP = ["Day", "Trial", "Direction", "Clock (EDT)", "Time (s)", "Segment (m)", "Segment (0-90 order)"] + list(COLS.values())
allr = allr[KEEP]
TRIALS = [("Day 1 (10/15)", "ESP-NOW walk"), ("Day 2 (10/20)", "Trial 1"), ("Day 2 (10/20)", "Trial 3"),
          ("Day 2 (10/20)", "Trial 4"), ("Day 2 (10/20)", "Return run")]
SHEET = {"ESP-NOW walk": "Day 1 ESP-NOW walk", "Trial 1": "Day 2 Trial 1", "Trial 3": "Day 2 Trial 3",
         "Trial 4": "Day 2 Trial 4", "Return run": "Day 2 Return run"}

# ---------------------------------------------------------------- segment summary
# Estimated position on the tape at the end of each segment. Radios measure distance to the base at 0 m;
# GPS/IMU measure distance travelled since zero, so on 90->0 runs position = 90 - value.
SENS = {"GPS": "GPS dist from zero (m)", "LoRa": "LoRa est dist (m)", "ESP-NOW": "ESP-NOW est dist (m)",
        "IMU1": "IMU1 dist from zero (m)", "IMU2": "IMU2 dist from zero (m)"}
rows = []
for day, tr in TRIALS:
    g = allr[allr.Trial == tr]
    for seg, gs in g.groupby("Segment (m)", sort=False):
        end = int(seg.split("-")[1])
        r = {"Day": day, "Trial": tr, "Direction": DIR[tr], "Segment (m)": seg, "True position (m)": end}
        gps = gs[SENS["GPS"]].dropna()
        r["Note"] = "No movement recorded; left out of graphs" if len(gps) and gps.max() - gps.min() < 1 else ""
        for name, col in SENS.items():
            v = gs[col].dropna()
            if not len(v):
                r[f"{name} est position (m)"] = r[f"{name} error (m)"] = None
                continue
            pos = 90 - v.iloc[-1] if DIR[tr] == "90->0" and name in ("GPS", "IMU1", "IMU2") else v.iloc[-1]
            r[f"{name} est position (m)"] = round(pos, 3)
            r[f"{name} error (m)"] = round(pos - end, 3)
        rows.append(r)
summ = pd.DataFrame(rows)

# ---------------------------------------------------------------- day 1 IMU walk (1 row per second)
imu["sec"] = imu.t_s.round()
imu1 = imu.groupby("sec").agg(walk_state=("walk_state", "last"), heading=("heading_deg", "mean"),
                              roll=("roll_deg", "mean"), pitch=("pitch_deg", "mean")).reset_index()
imu1.columns = ["Time (s)", "Walk state", "Heading (deg)", "Roll (deg)", "Pitch (deg)"]
imu1 = imu1.round(2)
walk = imu.dropna(subset=["integration_m"]).iloc[-1]
imu1.loc[imu1.index[-1], "IMU integrated distance (m)"] = round(walk.integration_m, 2)
imu1.loc[imu1.index[-1], "Walk duration (s)"] = walk.walk_duration_s

notes = pd.DataFrame({"Item": [
    "Sources", "Rows", "Segment (m)", "90->0 runs", "Left out", "Day 1 ESP-NOW walk", "Day 1 IMU walk",
    "Day 2 Trial 1", "Day 2 Trial 3", "Day 2 Trial 4", "Day 2 Return run"], "Note": [
    "Day 1: Downloads/Data.zip (ESP-NOW CSVs, IMU Data.csv). Day 2: Downloads/Data 2 (zips + loose overview CSVs). Originals untouched.",
    "Day 2 is one row per second from the dashboard overview logs (all sensors). Repeated rows across exports were de-duplicated.",
    "From the label typed in the dashboard. Rows logged while a label was being typed are assigned to that segment.",
    "Radios read distance to the base at 0 m. GPS/IMU read distance travelled from the 90 m end, so the summary uses 90 - value as position.",
    "Day 2 Trial 2 (aborted, only 0-10 walked). Rows logged before the sensors were zeroed at the start of a run. Raw 50 Hz IMU, packet and console logs.",
    "15:24-15:36 EDT, 0-80 m (no 80-90 file). ESP-NOW only. Duplicate (2) files merged. Link was lost at times from 30 m on.",
    "16:17 EDT, separate 88 s IMU walk (not on the segment timeline). IMU integrated 252 m; the true distance was not logged.",
    "13:43 EDT, 0->90 m. No zeroing between segments.",
    "14:33 EDT, 0->90 m. Starts at the re-zero after Trial 2.",
    "14:38 EDT, 90->0 m. The 10-0 segment shows no movement (GPS holds at 78.6 m), so it is left out of the graphs.",
    "13:52 EDT, 90->0 m after Trial 1. No trial number in the file names."]})

from openpyxl.styles import Font, PatternFill
with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    allr.to_excel(w, sheet_name="All trials 0-90", index=False)
    for day, tr in TRIALS:
        t = allr[allr.Trial == tr].drop(columns=["Day", "Trial", "Direction", "Segment (0-90 order)"]).dropna(axis=1, how="all")
        t.to_excel(w, sheet_name=SHEET[tr], index=False)
    imu1.to_excel(w, sheet_name="Day 1 IMU walk", index=False)
    summ.to_excel(w, sheet_name="Segment summary", index=False)
    notes.to_excel(w, sheet_name="Notes", index=False)
    for ws in w.book.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:300])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 110 if ws.title == "Notes" else 28)
allr.to_csv(OUT.replace(".xlsx", "_all_trials.csv"), index=False)
print(allr.groupby(["Day", "Trial"]).size().to_string())
print(f"Saved {OUT}")
